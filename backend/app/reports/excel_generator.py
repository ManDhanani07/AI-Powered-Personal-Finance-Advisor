"""
OpenPyXL Excel Generation Module for AI-Powered Personal Finance Advisor.
Produces clean, beautifully styled spreadsheets with high readability, formatted currency,
auto-fitted column widths, frozen headers, and multiple organized worksheets.
"""

import io
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Styling Palettes
FONT_FAMILY = "Segoe UI"
COLOR_PRIMARY_DARK = "0F172A"   # Slate 900
COLOR_HEADER_NAVY = "1E293B"    # Slate 800
COLOR_ACCENT_EMERALD = "10B981" # Emerald 500
COLOR_ACCENT_BG = "F0FDF4"      # Emerald 50
COLOR_BG_SUBTLE = "F8FAFC"      # Slate 50
COLOR_BORDER = "CBD5E1"         # Slate 300
COLOR_BORDER_LIGHT = "E2E8F0"   # Slate 200

# Fonts
FONT_TITLE = Font(name=FONT_FAMILY, size=15, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name=FONT_FAMILY, size=9, italic=True, color="CBD5E1")
FONT_SECTION = Font(name=FONT_FAMILY, size=11, bold=True, color="0F172A")
FONT_TH = Font(name=FONT_FAMILY, size=9.5, bold=True, color="FFFFFF")
FONT_CELL = Font(name=FONT_FAMILY, size=9.5, color="1E293B")
FONT_CELL_BOLD = Font(name=FONT_FAMILY, size=9.5, bold=True, color="1E293B")
FONT_CELL_MUTED = Font(name=FONT_FAMILY, size=9, italic=True, color="64748B")

# Fills
FILL_PRIMARY = PatternFill(start_color=COLOR_PRIMARY_DARK, end_color=COLOR_PRIMARY_DARK, fill_type="solid")
FILL_HEADER = PatternFill(start_color=COLOR_HEADER_NAVY, end_color=COLOR_HEADER_NAVY, fill_type="solid")
FILL_ACCENT = PatternFill(start_color=COLOR_ACCENT_BG, end_color=COLOR_ACCENT_BG, fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=COLOR_BG_SUBTLE, end_color=COLOR_BG_SUBTLE, fill_type="solid")
FILL_TOTAL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

# Borders
BORDER_THIN = Border(
    left=Side(style="thin", color=COLOR_BORDER_LIGHT),
    right=Side(style="thin", color=COLOR_BORDER_LIGHT),
    top=Side(style="thin", color=COLOR_BORDER_LIGHT),
    bottom=Side(style="thin", color=COLOR_BORDER_LIGHT),
)
BORDER_TOP_BOTTOM_DOUBLE = Border(
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="double", color="0F172A"),
)

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_BANNER = Alignment(horizontal="left", vertical="center", indent=1)

# Number Formats
FMT_CURRENCY = '"₹"#,##0.00'
FMT_PERCENT = '0.0%'
FMT_INTEGER = '#,##0'
FMT_DATE = 'yyyy-mm-dd'


def _autofit_columns(ws, min_width=12, max_width=45):
    """Dynamically calculates maximum string length in each column and applies proper padding."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            # Don't let merged header banner inflate column width calculation
            if cell.row in [1, 2]:
                continue
            if "\n" in val_str:
                lines = val_str.split("\n")
                line_max = max(len(l) for l in lines)
                max_len = max(max_len, line_max)
            else:
                max_len = max(max_len, len(val_str))

        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 5, max_width))


def generate_excel_financial_report(
    user_display: str,
    user_email: str,
    filter_type: str,
    summary_rep: Dict[str, Any],
    cat_rep: Dict[str, Any],
    budget_rep: Dict[str, Any],
    goal_rep: Dict[str, Any],
    health_rep: Dict[str, Any],
    forecast_rep: Dict[str, Any],
    transactions: List[Any],
    report_type: str = "executive",
) -> io.BytesIO:
    """
    Builds an executive-ready multi-tab Excel workbook with formatted cells,
    currency numbers, auto-filters, and clean typography.
    """
    wb = openpyxl.Workbook()
    now_str = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")

    # ─────────────────────────────────────────────────────────────
    # TAB 1: EXECUTIVE SUMMARY
    # ─────────────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Executive Summary"
    ws_sum.views.sheetView[0].showGridLines = True

    # 1. Header Banner
    ws_sum.merge_cells("A1:E1")
    title_cell = ws_sum["A1"]
    title_cell.value = "AI WEALTH OS - FINANCIAL EXECUTIVE REPORT"
    title_cell.font = FONT_TITLE
    title_cell.fill = FILL_PRIMARY
    title_cell.alignment = ALIGN_BANNER
    ws_sum.row_dimensions[1].height = 36

    ws_sum.merge_cells("A2:E2")
    sub_cell = ws_sum["A2"]
    sub_cell.value = f"Account: {user_display} ({user_email}) | Filter Scope: {filter_type.upper()} | Generated: {now_str}"
    sub_cell.font = FONT_SUBTITLE
    sub_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_cell.alignment = ALIGN_BANNER
    ws_sum.row_dimensions[2].height = 20

    ws_sum.append([]) # Spacer row 3

    # 2. Key Metrics Section
    ws_sum.append(["EXECUTIVE FINANCIAL METRICS", "", "", "", ""])
    ws_sum.merge_cells("A4:E4")
    ws_sum["A4"].font = FONT_SECTION
    ws_sum.row_dimensions[4].height = 24

    ws_sum.append(["Metric Name", "Value", "Benchmark / Status", "Category", "Notes"])
    for col_idx in range(1, 6):
        c = ws_sum.cell(row=5, column=col_idx)
        c.font = FONT_TH
        c.fill = FILL_HEADER
        c.alignment = ALIGN_LEFT if col_idx != 2 else ALIGN_RIGHT
    ws_sum.row_dimensions[5].height = 24

    kpis = summary_rep.get("kpis", {})
    inc_val = float(kpis.get("total_income", 0.0))
    exp_val = float(kpis.get("total_expenses", 0.0))
    sav_val = float(kpis.get("net_savings", 0.0))
    sav_rate = float(kpis.get("savings_rate", 0.0))
    health_score = float(health_rep.get("latest_score", 75.0))
    health_grade = health_rep.get("latest_grade", "B")

    metrics_rows = [
        ("Total Inflow (Income)", inc_val, "Gross Inflow", "Cashflow", "Aggregated salary, returns & credits", FMT_CURRENCY),
        ("Total Outflow (Expenses)", exp_val, "Gross Outflow", "Cashflow", "Living, routine, and discretionary expenses", FMT_CURRENCY),
        ("Net Savings / Surplus", sav_val, "Active Surplus" if sav_val >= 0 else "Deficit", "Cashflow", "Surplus cash added to savings cushion", FMT_CURRENCY),
        ("Savings Rate", sav_rate / 100.0, "Healthy" if sav_rate >= 20 else "Fair", "Efficiency", "Proportion of total income conserved", FMT_PERCENT),
        ("Financial Health Score", health_score, f"Grade {health_grade}", "Diagnostic", "Weighted 4-pillar financial audit score", FMT_INTEGER),
        ("ML Expense Forecast (Next Mo)", float(forecast_rep.get("forecast_next_month_expense", 0.0)), "Reliability: " + str(forecast_rep.get("forecast_reliability", "Moderate")), "Predictive", "Gradient Boosted LightGBM Master Forecast", FMT_CURRENCY),
    ]

    for r_idx, (name, val, status, cat, notes, num_fmt) in enumerate(metrics_rows, start=6):
        row_cells = [
            ws_sum.cell(row=r_idx, column=1, value=name),
            ws_sum.cell(row=r_idx, column=2, value=val),
            ws_sum.cell(row=r_idx, column=3, value=status),
            ws_sum.cell(row=r_idx, column=4, value=cat),
            ws_sum.cell(row=r_idx, column=5, value=notes),
        ]
        row_fill = FILL_ZEBRA if r_idx % 2 == 1 else PatternFill(fill_type=None)
        for c_i, cell in enumerate(row_cells, start=1):
            cell.font = FONT_CELL_BOLD if c_i in [1, 2] else FONT_CELL
            cell.border = BORDER_THIN
            if row_fill.fill_type:
                cell.fill = row_fill
            if c_i == 2:
                cell.number_format = num_fmt
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT
        ws_sum.row_dimensions[r_idx].height = 20

    # 3. AI Strategic Financial Assessment
    next_r = len(metrics_rows) + 7
    ws_sum.cell(row=next_r, column=1, value="AI COPILOT STRATEGIC ADVISORY ASSESSMENT").font = FONT_SECTION
    ws_sum.merge_cells(f"A{next_r}:E{next_r}")
    ws_sum.row_dimensions[next_r].height = 22

    exec_insight = summary_rep.get("executive_insight") or "Financial cash flow position remains balanced. Routine expenses are well within safe thresholds."
    health_expl = health_rep.get("explanation") or "Composite financial score indicates strong budgetary discipline and resilience."

    advice_r1 = next_r + 1
    ws_sum.cell(row=advice_r1, column=1, value="Executive Insight:")
    ws_sum.cell(row=advice_r1, column=1).font = FONT_CELL_BOLD
    ws_sum.cell(row=advice_r1, column=2, value=exec_insight)
    ws_sum.merge_cells(f"B{advice_r1}:E{advice_r1}")
    ws_sum.row_dimensions[advice_r1].height = 28

    advice_r2 = next_r + 2
    ws_sum.cell(row=advice_r2, column=1, value="Health Audit:")
    ws_sum.cell(row=advice_r2, column=1).font = FONT_CELL_BOLD
    ws_sum.cell(row=advice_r2, column=2, value=health_expl)
    ws_sum.merge_cells(f"B{advice_r2}:E{advice_r2}")
    ws_sum.row_dimensions[advice_r2].height = 28

    _autofit_columns(ws_sum)

    # ─────────────────────────────────────────────────────────────
    # TAB 2: CATEGORY SPENDING
    # ─────────────────────────────────────────────────────────────
    ws_cat = wb.create_sheet(title="Category Spending")
    ws_cat.views.sheetView[0].showGridLines = True

    ws_cat.merge_cells("A1:E1")
    cat_banner = ws_cat["A1"]
    cat_banner.value = "CATEGORY EXPENDITURE BREAKDOWN & BUDGET ALLOCATION"
    cat_banner.font = FONT_TITLE
    cat_banner.fill = FILL_PRIMARY
    cat_banner.alignment = ALIGN_BANNER
    ws_cat.row_dimensions[1].height = 36

    ws_cat.append([])
    ws_cat.append(["Category Name", "Total Amount", "Share of Spend", "Transactions Count", "Average Ticket Size"])
    for col_idx in range(1, 6):
        c = ws_cat.cell(row=3, column=col_idx)
        c.font = FONT_TH
        c.fill = FILL_HEADER
        c.alignment = ALIGN_LEFT if col_idx == 1 else ALIGN_RIGHT
    ws_cat.row_dimensions[3].height = 24
    ws_cat.freeze_panes = "A4"

    cat_items = cat_rep.get("categories", [])
    tot_cat_spend = sum(float(c.get("amount", 0.0)) for c in cat_items)

    r_start = 4
    if cat_items:
        for idx, c in enumerate(cat_items, start=r_start):
            c_name = str(c.get("category_name") or c.get("name") or "Uncategorized").title()
            c_amt = float(c.get("amount", 0.0))
            c_pct = (c_amt / tot_cat_spend) if tot_cat_spend > 0 else 0.0
            c_count = int(c.get("count", 0) or c.get("transaction_count", 0))
            avg_ticket = (c_amt / c_count) if c_count > 0 else 0.0

            ws_cat.cell(row=idx, column=1, value=c_name).alignment = ALIGN_LEFT
            
            c2 = ws_cat.cell(row=idx, column=2, value=c_amt)
            c2.number_format = FMT_CURRENCY
            c2.alignment = ALIGN_RIGHT

            c3 = ws_cat.cell(row=idx, column=3, value=c_pct)
            c3.number_format = FMT_PERCENT
            c3.alignment = ALIGN_RIGHT

            c4 = ws_cat.cell(row=idx, column=4, value=c_count)
            c4.number_format = FMT_INTEGER
            c4.alignment = ALIGN_RIGHT

            c5 = ws_cat.cell(row=idx, column=5, value=avg_ticket)
            c5.number_format = FMT_CURRENCY
            c5.alignment = ALIGN_RIGHT

            row_fill = FILL_ZEBRA if (idx % 2 == 1) else PatternFill(fill_type=None)
            for c_i in range(1, 6):
                cell = ws_cat.cell(row=idx, column=c_i)
                cell.font = FONT_CELL
                cell.border = BORDER_THIN
                if row_fill.fill_type:
                    cell.fill = row_fill
            ws_cat.row_dimensions[idx].height = 20

        # Totals Row
        tot_r = len(cat_items) + r_start
        ws_cat.cell(row=tot_r, column=1, value="TOTAL CATEGORIZED EXPENDITURE").font = FONT_CELL_BOLD
        
        tot_c2 = ws_cat.cell(row=tot_r, column=2, value=tot_cat_spend)
        tot_c2.font = FONT_CELL_BOLD
        tot_c2.number_format = FMT_CURRENCY
        tot_c2.alignment = ALIGN_RIGHT

        tot_c3 = ws_cat.cell(row=tot_r, column=3, value=1.0)
        tot_c3.font = FONT_CELL_BOLD
        tot_c3.number_format = FMT_PERCENT
        tot_c3.alignment = ALIGN_RIGHT

        tot_c4 = ws_cat.cell(row=tot_r, column=4, value=sum(int(c.get("count", 0) or c.get("transaction_count", 0)) for c in cat_items))
        tot_c4.font = FONT_CELL_BOLD
        tot_c4.number_format = FMT_INTEGER
        tot_c4.alignment = ALIGN_RIGHT

        tot_c5 = ws_cat.cell(row=tot_r, column=5, value=(tot_cat_spend / max(1, tot_c4.value)))
        tot_c5.font = FONT_CELL_BOLD
        tot_c5.number_format = FMT_CURRENCY
        tot_c5.alignment = ALIGN_RIGHT

        for c_i in range(1, 6):
            cell = ws_cat.cell(row=tot_r, column=c_i)
            cell.fill = FILL_TOTAL
            cell.border = BORDER_TOP_BOTTOM_DOUBLE
        ws_cat.row_dimensions[tot_r].height = 22
    else:
        ws_cat.append(["No category data logged for selected period.", 0, 0, 0, 0])

    _autofit_columns(ws_cat)

    # ─────────────────────────────────────────────────────────────
    # TAB 3: BUDGETS & GOALS
    # ─────────────────────────────────────────────────────────────
    ws_bg = wb.create_sheet(title="Budgets & Goals")
    ws_bg.views.sheetView[0].showGridLines = True

    ws_bg.merge_cells("A1:E1")
    bg_banner = ws_bg["A1"]
    bg_banner.value = "BUDGET UTILIZATION & FINANCIAL GOALS AUDIT"
    bg_banner.font = FONT_TITLE
    bg_banner.fill = FILL_PRIMARY
    bg_banner.alignment = ALIGN_BANNER
    ws_bg.row_dimensions[1].height = 36

    ws_bg.append([])
    ws_bg.append(["BUDGET CATEGORY CAPS", "", "", "", ""])
    ws_bg.merge_cells("A3:E3")
    ws_bg["A3"].font = FONT_SECTION
    ws_bg.row_dimensions[3].height = 22

    ws_bg.append(["Category / Scope", "Configured Limit", "Spent Amount", "Utilization", "Current Status"])
    for col_idx in range(1, 6):
        c = ws_bg.cell(row=4, column=col_idx)
        c.font = FONT_TH
        c.fill = FILL_HEADER
        c.alignment = ALIGN_LEFT if col_idx in [1, 5] else ALIGN_RIGHT
    ws_bg.row_dimensions[4].height = 22

    b_limit = float(budget_rep.get("total_limit", 0.0))
    b_spent = float(budget_rep.get("total_spent", 0.0))
    b_util = float(budget_rep.get("overall_utilization_pct", 0.0))
    b_status = "Healthy" if b_util <= 85 else "Approaching Limit"

    ws_bg.append(["Overall Monthly Budget Cap", b_limit, b_spent, b_util / 100.0, b_status])
    b_row = 5
    for c_i in range(1, 6):
        cell = ws_bg.cell(row=b_row, column=c_i)
        cell.font = FONT_CELL_BOLD if c_i == 1 else FONT_CELL
        cell.border = BORDER_THIN
        if c_i in [2, 3]:
            cell.number_format = FMT_CURRENCY
            cell.alignment = ALIGN_RIGHT
        elif c_i == 4:
            cell.number_format = FMT_PERCENT
            cell.alignment = ALIGN_RIGHT
        else:
            cell.alignment = ALIGN_LEFT

    ws_bg.append([])
    g_title_row = 7
    ws_bg.cell(row=g_title_row, column=1, value="FINANCIAL WEALTH GOALS PROGRESS").font = FONT_SECTION
    ws_bg.merge_cells(f"A{g_title_row}:E{g_title_row}")
    ws_bg.row_dimensions[g_title_row].height = 22

    ws_bg.append(["Goal Name", "Target Amount", "Current Saved", "Completion %", "Timeline / Status"])
    g_header_row = 8
    for col_idx in range(1, 6):
        c = ws_bg.cell(row=g_header_row, column=col_idx)
        c.font = FONT_TH
        c.fill = FILL_HEADER
        c.alignment = ALIGN_LEFT if col_idx in [1, 5] else ALIGN_RIGHT
    ws_bg.row_dimensions[g_header_row].height = 22

    goals_list = goal_rep.get("goals", [])
    if goals_list:
        for g_idx, g in enumerate(goals_list, start=9):
            g_name = str(g.get("name") or g.get("title") or "Savings Goal").title()
            g_target = float(g.get("target_amount", 0.0))
            g_saved = float(g.get("current_amount", 0.0) or g.get("saved_amount", 0.0))
            g_pct = (g_saved / g_target) if g_target > 0 else 0.0
            g_stat = str(g.get("status", "In Progress")).title()

            ws_bg.cell(row=g_idx, column=1, value=g_name).alignment = ALIGN_LEFT
            
            c2 = ws_bg.cell(row=g_idx, column=2, value=g_target)
            c2.number_format = FMT_CURRENCY
            c2.alignment = ALIGN_RIGHT

            c3 = ws_bg.cell(row=g_idx, column=3, value=g_saved)
            c3.number_format = FMT_CURRENCY
            c3.alignment = ALIGN_RIGHT

            c4 = ws_bg.cell(row=g_idx, column=4, value=g_pct)
            c4.number_format = FMT_PERCENT
            c4.alignment = ALIGN_RIGHT

            ws_bg.cell(row=g_idx, column=5, value=g_stat).alignment = ALIGN_LEFT

            for c_i in range(1, 6):
                ws_bg.cell(row=g_idx, column=c_i).font = FONT_CELL
                ws_bg.cell(row=g_idx, column=c_i).border = BORDER_THIN
            ws_bg.row_dimensions[g_idx].height = 20
    else:
        # Default entry
        g_total = int(goal_rep.get("total_goals", 0))
        g_pct = float(goal_rep.get("overall_completion_pct", 0.0))
        ws_bg.append([f"Consolidated Portfolio Goals ({g_total} active)", 100000.0, 100000.0 * (g_pct / 100.0), g_pct / 100.0, "Active"])
        r_def = 9
        for c_i in range(1, 6):
            cell = ws_bg.cell(row=r_def, column=c_i)
            cell.font = FONT_CELL
            cell.border = BORDER_THIN
            if c_i in [2, 3]:
                cell.number_format = FMT_CURRENCY
                cell.alignment = ALIGN_RIGHT
            elif c_i == 4:
                cell.number_format = FMT_PERCENT
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT

    _autofit_columns(ws_bg)

    # ─────────────────────────────────────────────────────────────
    # TAB 4: DETAILED TRANSACTIONS LEDGER
    # ─────────────────────────────────────────────────────────────
    ws_tx = wb.create_sheet(title="Audited Transactions")
    ws_tx.views.sheetView[0].showGridLines = True

    ws_tx.merge_cells("A1:G1")
    tx_banner = ws_tx["A1"]
    tx_banner.value = "AUDITED FINANCIAL TRANSACTIONS LEDGER"
    tx_banner.font = FONT_TITLE
    tx_banner.fill = FILL_PRIMARY
    tx_banner.alignment = ALIGN_BANNER
    ws_tx.row_dimensions[1].height = 36

    ws_tx.append([])
    headers = ["Date", "Title / Description", "Category", "Transaction Type", "Amount", "Payment Method", "Merchant"]
    ws_tx.append(headers)
    for col_idx in range(1, 8):
        c = ws_tx.cell(row=3, column=col_idx)
        c.font = FONT_TH
        c.fill = FILL_HEADER
        c.alignment = ALIGN_LEFT if col_idx != 5 else ALIGN_RIGHT
    ws_tx.row_dimensions[3].height = 24
    ws_tx.freeze_panes = "A4"

    for idx, t in enumerate(transactions, start=4):
        t_date = t.transaction_date.strftime("%Y-%m-%d") if hasattr(t, "transaction_date") and t.transaction_date else "N/A"
        t_title = getattr(t, "title", "Transaction") or "Transaction"
        t_cat = (t.category.category_name if hasattr(t, "category") and t.category else "Uncategorized") or "Uncategorized"
        t_type = (getattr(t, "transaction_type", "EXPENSE") or "EXPENSE").upper()
        t_amt = float(getattr(t, "amount", 0.0))
        t_method = getattr(t, "payment_method", "UPI") or "UPI"
        t_merchant = getattr(t, "merchant", "General") or "N/A"

        ws_tx.cell(row=idx, column=1, value=t_date).alignment = ALIGN_LEFT
        ws_tx.cell(row=idx, column=2, value=t_title).alignment = ALIGN_LEFT
        ws_tx.cell(row=idx, column=3, value=t_cat).alignment = ALIGN_LEFT
        ws_tx.cell(row=idx, column=4, value=t_type).alignment = ALIGN_LEFT
        
        amt_cell = ws_tx.cell(row=idx, column=5, value=t_amt)
        amt_cell.number_format = FMT_CURRENCY
        amt_cell.alignment = ALIGN_RIGHT

        ws_tx.cell(row=idx, column=6, value=t_method).alignment = ALIGN_LEFT
        ws_tx.cell(row=idx, column=7, value=t_merchant).alignment = ALIGN_LEFT

        row_fill = FILL_ZEBRA if (idx % 2 == 1) else PatternFill(fill_type=None)
        for c_i in range(1, 8):
            cell = ws_tx.cell(row=idx, column=c_i)
            cell.font = FONT_CELL
            cell.border = BORDER_THIN
            if row_fill.fill_type:
                cell.fill = row_fill
        ws_tx.row_dimensions[idx].height = 20

    if len(transactions) > 0:
        # Enable Excel Auto-Filter on data rows
        ws_tx.auto_filter.ref = f"A3:G{len(transactions) + 3}"

    _autofit_columns(ws_tx)

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
